"""Render one knowledge object as Markdown, Word, or PDF."""
from __future__ import annotations

import io
import json
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


KIND_LABELS = {
    "concept": "概念", "resource": "资料", "repository": "代码仓库", "model": "模型",
    "dataset": "数据集", "benchmark": "评测", "experiment": "实验", "question": "问题", "task": "任务",
}
RELATION_LABELS = {
    "depends_on": "依赖", "implements": "实现", "uses": "使用", "evaluates": "评测",
    "supports": "支持", "questions": "质疑", "compares_with": "比较", "related_to": "相关",
}


def material_markdown(record: dict) -> str:
    obj = record["object"]
    lines = [f"# {obj['title']}", "", f"- 类型：{KIND_LABELS.get(obj['kind'], obj['kind'])}"]
    if obj.get("tags"):
        lines.append(f"- 标签：{', '.join(obj['tags'])}")
    if obj.get("source_url"):
        lines.append(f"- 来源：<{obj['source_url']}>")
    lines.append(f"- 导出版本：{obj.get('version', 1)}")
    if obj.get("summary"):
        lines.extend(["", "## 摘要", "", obj["summary"]])
    lines.extend(["", "## 正文", "", obj.get("body_md") or "（暂无正文）"])
    if record.get("relations"):
        lines.extend(["", "## 相关知识与证据", ""])
        for relation in record["relations"]:
            direction = "→" if str(relation["source_id"]) == str(obj["id"]) else "←"
            other = relation["target_title"] if direction == "→" else relation["source_title"]
            lines.append(f"### {direction} {RELATION_LABELS.get(relation['relation_type'], relation['relation_type'])}：{other}")
            if relation.get("description"):
                lines.extend(["", relation["description"]])
            for evidence in relation.get("evidence", []):
                detail = evidence.get("explanation") or "关系证据"
                locator = f"（{evidence['locator']}）" if evidence.get("locator") else ""
                source = evidence.get("source_url")
                lines.append(f"- {detail}{locator}" + (f"：<{source}>" if source else ""))
            lines.append("")
    if record.get("attachments"):
        lines.extend(["## 附件", ""])
        for item in record["attachments"]:
            lines.append(f"- {item['original_name']}（{item['size_bytes']} bytes）")
        lines.append("")
    if obj.get("metadata"):
        lines.extend(["## 扩展属性", "", "```json", json.dumps(obj["metadata"], ensure_ascii=False, indent=2), "```", ""])
    return "\n".join(lines).rstrip() + "\n"


def _markdown_lines(markdown: str):
    in_code = False
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            yield "code", line
        elif match := re.match(r"^(#{1,4})\s+(.*)$", line):
            yield f"h{len(match.group(1))}", match.group(2)
        elif match := re.match(r"^[-*]\s+(.*)$", line):
            yield "bullet", match.group(1)
        elif match := re.match(r"^\d+[.)]\s+(.*)$", line):
            yield "number", match.group(1)
        elif line.startswith("> "):
            yield "quote", line[2:]
        elif not line:
            yield "blank", ""
        else:
            yield "paragraph", line


def material_docx(markdown: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    for kind, text in _markdown_lines(markdown):
        if kind.startswith("h"):
            document.add_heading(text, level=min(int(kind[1:]), 4))
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        elif kind == "number":
            document.add_paragraph(text, style="List Number")
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif kind == "quote":
            document.add_paragraph(text, style="Quote")
        elif kind == "paragraph":
            document.add_paragraph(text)
        elif kind == "blank" and document.paragraphs and document.paragraphs[-1].text:
            document.add_paragraph()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def material_pdf(markdown: str) -> bytes:
    try:
        pdfmetrics.getFont("STSong-Light")
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm)
    base = getSampleStyleSheet()
    styles = {
        "h1": ParagraphStyle("CJK-H1", parent=base["Title"], fontName="STSong-Light", fontSize=22, leading=30, alignment=TA_CENTER, spaceAfter=14),
        "h2": ParagraphStyle("CJK-H2", parent=base["Heading2"], fontName="STSong-Light", fontSize=16, leading=22, spaceBefore=12, spaceAfter=7),
        "h3": ParagraphStyle("CJK-H3", parent=base["Heading3"], fontName="STSong-Light", fontSize=13, leading=19, spaceBefore=9, spaceAfter=5),
        "h4": ParagraphStyle("CJK-H4", parent=base["Heading4"], fontName="STSong-Light", fontSize=11, leading=17, spaceBefore=7),
        "paragraph": ParagraphStyle("CJK-Body", parent=base["BodyText"], fontName="STSong-Light", fontSize=9.5, leading=16, wordWrap="CJK", spaceAfter=5),
        "bullet": ParagraphStyle("CJK-Bullet", parent=base["BodyText"], fontName="STSong-Light", fontSize=9.5, leading=16, leftIndent=12, firstLineIndent=-8, wordWrap="CJK"),
        "number": ParagraphStyle("CJK-Number", parent=base["BodyText"], fontName="STSong-Light", fontSize=9.5, leading=16, leftIndent=12, wordWrap="CJK"),
        "quote": ParagraphStyle("CJK-Quote", parent=base["BodyText"], fontName="STSong-Light", fontSize=9.5, leading=16, leftIndent=12, textColor="#58735a", wordWrap="CJK"),
        "code": ParagraphStyle("CJK-Code", parent=base["Code"], fontName="STSong-Light", fontSize=8, leading=12, leftIndent=8, wordWrap="CJK", backColor="#f3f5ef"),
    }
    story = []
    for kind, text in _markdown_lines(markdown):
        if kind == "blank":
            story.append(Spacer(1, 3 * mm))
            continue
        value = escape(text).replace("  ", " &nbsp;")
        prefix = "• " if kind == "bullet" else ""
        story.append(Paragraph(prefix + value, styles.get(kind, styles["paragraph"])))
    document.build(story)
    return output.getvalue()
